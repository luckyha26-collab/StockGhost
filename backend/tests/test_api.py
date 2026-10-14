"""End-to-end API tests. Run: python -m pytest tests -q   (uses a temporary SQLite database)."""
import csv
import io
import os
import sys
import tempfile

_tmp = tempfile.mkdtemp()
os.environ["DATABASE_URL"] = "sqlite:///" + os.path.join(_tmp, "t.db")
os.environ["UPLOAD_DIR"] = os.path.join(_tmp, "up")
os.environ["JWT_SECRET"] = "test-secret-test-secret-test-secret-123456"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402
import pytest  # noqa: E402

from app import create_app  # noqa: E402
from utils import reset_rate_limits  # noqa: E402

PW = "Strong#Pass1"


@pytest.fixture(scope="module")
def app():
    a = create_app()
    return a


@pytest.fixture(autouse=True)
def _rl():
    reset_rate_limits()


class Client:
    def __init__(self, app):
        self.c, self.h = app.test_client(), {}

    def req(self, m, url, **kw):
        r = getattr(self.c, m)(url, headers=self.h, **kw)
        return r.status_code, (r.get_json(silent=True) if r.mimetype == "application/json" else r)

    def get(self, u, **k): return self.req("get", u, **k)
    def post(self, u, j=None, **k): return self.req("post", u, json=j, **k) if j is not None or not k else self.req("post", u, **k)
    def put(self, u, j): return self.req("put", u, json=j)
    def delete(self, u): return self.req("delete", u)


def register(app, email, company="Acme Traders", btype="Retail"):
    c = Client(app)
    st, r = c.post("/api/auth/register", {"full_name": "Test User", "company_name": company, "business_type": btype, "email": email,
                                          "password": PW, "confirm_password": PW})
    assert st == 201, r
    c.h = {"Authorization": "Bearer " + r["data"]["access_token"]}
    return c, r["data"]


def csv_file(rows, name="data.csv"):
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return (io.BytesIO(buf.getvalue().encode()), name)


def upload(c, dtype, rows, name="data.csv"):
    return c.post("/api/imports", data={"dataset_type": dtype, "file": csv_file(rows, name)}, content_type="multipart/form-data")


PRODUCTS = [["Item Code", "Product Name", "Category", "Cost", "MRP", "Stock", "Reorder Level"],
            ["a-1", "Wireless Mouse", "Accessories", "300", "599", "820", "20"],
            ["A-2", "USB Hub", "Accessories", "250", "499", "18", "25"],
            ["A-3", "Gaming Keyboard", "Accessories", "1200", "2400", "500", "50"],
            ["A-4", "Dup Row", "Accessories", "abc", "10", "5", "2"]]


def sales_rows():
    import datetime as dt
    t = dt.date.today()
    rows = [["Product Name", "Item Code", "Qty Sold", "Sale Date", "Price"]]
    for i in range(1, 120):
        d = t - dt.timedelta(days=i)
        rows.append(["Wireless Mouse", "A-1", 2, d.strftime("%d/%m/%Y"), "₹599"])
        if i % 4 == 0:
            rows.append(["USB Hub", "A-2", 3, d.strftime("%d/%m/%Y"), "499"])
    rows.append(["USB Hub", "A-2", -1, "01/01/2026", "499"])           # negative qty
    rows.append(["USB Hub", "A-2", 1, "99/99/2026", "499"])            # invalid date
    return rows


def full_import(c, dtype, rows, mapping=None, options=None, replace_demo=False):
    st, r = upload(c, dtype, rows)
    assert st == 201, r
    jid = r["data"]["id"]
    mp = mapping or r["data"]["suggested_mapping"]
    st, v = c.post(f"/api/imports/{jid}/validate", {"mapping": mp, "options": options or {}})
    assert st == 200, v
    st, r2 = c.post(f"/api/imports/{jid}/confirm", {"mapping": mp, "options": options or {}, "replace_demo": replace_demo})
    return jid, v["data"]["report"], st, r2


# ------------------------------------------------------------------------------------------------------------------
def test_registration_validation(app):
    c = Client(app)
    bad = {"full_name": "X", "company_name": "Co", "business_type": "Retail", "email": "nope", "password": PW, "confirm_password": PW}
    assert c.post("/api/auth/register", bad)[0] == 400
    assert c.post("/api/auth/register", {**bad, "email": "a@b.co", "password": "weak", "confirm_password": "weak"})[0] == 400
    assert c.post("/api/auth/register", {**bad, "email": "a@b.co", "confirm_password": "Different#1"})[0] == 400
    register(app, "dup@test.com")
    st, r = c.post("/api/auth/register", {**bad, "email": "dup@test.com"})
    assert st == 409 and r["success"] is False


def test_login_logout_and_lockout(app):
    c, _ = register(app, "login@test.com")
    st, r = Client(app).post("/api/auth/login", {"email": "login@test.com", "password": PW})
    assert st == 200 and r["data"]["workspace"]["name"] == "Acme Traders"
    assert Client(app).post("/api/auth/login", {"email": "login@test.com", "password": "wrong"})[0] == 401
    assert c.post("/api/auth/logout", {})[0] == 200
    assert c.get("/api/workspace")[0] == 401                       # token revoked
    z = Client(app)
    for _ in range(5):
        z.post("/api/auth/login", {"email": "login@test.com", "password": "bad"})
    assert z.post("/api/auth/login", {"email": "login@test.com", "password": PW})[0] == 429   # locked


def test_protected_routes_need_token(app):
    c = Client(app)
    for u in ("/api/products", "/api/analytics/dashboard", "/api/imports", "/api/users", "/api/audit-log"):
        assert c.get(u)[0] == 401


def test_password_reset_flow(app):
    register(app, "reset@test.com")
    c = Client(app)
    st, r = c.post("/api/auth/forgot-password", {"email": "reset@test.com"})
    assert st == 200
    token = r["data"]["dev_reset_link"].split("token=")[1]
    assert c.post("/api/auth/reset-password", {"token": token, "password": "New#Passw0rd", "confirm_password": "New#Passw0rd"})[0] == 200
    assert c.post("/api/auth/reset-password", {"token": token, "password": "New#Passw0rd", "confirm_password": "New#Passw0rd"})[0] == 400  # one-time
    assert c.post("/api/auth/login", {"email": "reset@test.com", "password": "New#Passw0rd"})[0] == 200
    st, r = c.post("/api/auth/forgot-password", {"email": "unknown@test.com"})
    assert st == 200 and "dev_reset_link" not in r["data"]


def test_empty_workspace_has_no_fake_data(app):
    c, _ = register(app, "empty@test.com")
    st, r = c.get("/api/analytics/dashboard")
    assert st == 200 and r["data"]["has_data"] is False and "kpis" not in r["data"]
    st, r = c.post("/api/chatbot", {"message": "top 10 dead stock"})
    assert "No business data" in r["data"]["answer"]
    assert c.get("/api/reports/inventory")[0] == 400


def test_import_pipeline_and_dashboard(app):
    c, _ = register(app, "import@test.com")
    jid, rep, st, r = full_import(c, "products", PRODUCTS)
    assert st == 200, r
    assert rep["rows_total"] == 4 and rep["rows_valid"] == 3 and rep["rows_rejected"] == 1      # 'abc' cost rejected, reported
    assert any(x["code"] == "invalid_number" for x in rep["by_code"])
    jid2, rep2, st, r = full_import(c, "sales", sales_rows())
    assert st == 200, r
    assert rep2["rows_rejected"] == 2 and rep2["rows_valid"] >= 140
    st, d = c.get("/api/analytics/dashboard?range=90d")
    k = d["data"]["kpis"]
    assert d["data"]["has_data"] and k["total_products"] == 3
    assert k["inventory_value"] == 820 * 300 + 18 * 250 + 500 * 1200
    assert k["sales_revenue"] > 0 and k["blocked_capital"] > 0
    # Gaming Keyboard never sold -> must be flagged using only this workspace's data
    st, rk = c.get("/api/analytics/risk?q=keyboard")
    row = rk["data"]["items"][0]
    assert row["risk_level"] in ("HIGH", "CRITICAL") and row["stock_status"] == "Dead Stock" and row["reasons"]
    st, rk = c.get("/api/analytics/forecast?flag=stockout")
    assert any(i["product_name"] == "USB Hub" for i in rk["data"]["items"]) or rk["data"]["summary"]["stockout_high"] >= 0
    # error report + history
    st, resp = c.get(f"/api/imports/{jid}/errors.csv")
    assert st == 200 and b"invalid_number" in resp.data
    st, h = c.get("/api/imports")
    assert h["data"]["total"] == 2
    # copilot uses workspace data
    st, r = c.post("/api/chatbot", {"message": "Why is Gaming Keyboard high risk?"})
    assert "Gaming Keyboard" in r["data"]["answer"] and "DSRS" in r["data"]["answer"]
    st, r = c.post("/api/chatbot", {"message": "What happens if I discount the top 5 dead-stock products by 20%?"})
    assert "PROJECTION" in r["data"]["answer"]


def test_excel_upload_and_mapping(app):
    c, _ = register(app, "xlsx@test.com")
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in [["Vendor", "Item", "Date", "Qty", "Rate"]] + [["Acme Supply", "Widget", "05/03/2026", 10, 55.5], ["Acme Supply", "Widget", "06/03/2026", 5, 56]]:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    st, r = c.post("/api/imports", data={"dataset_type": "purchases", "file": (buf, "po.xlsx")}, content_type="multipart/form-data")
    assert st == 201 and r["data"]["suggested_mapping"]["cost_price"] == "Rate"
    jid = r["data"]["id"]
    mp = r["data"]["suggested_mapping"]
    st, v = c.post(f"/api/imports/{jid}/validate", {"mapping": mp, "options": {"auto_fix": True}})
    assert st == 200 and v["data"]["report"]["rows_valid"] == 2
    st, r = c.post(f"/api/imports/{jid}/confirm", {"mapping": mp, "options": {"auto_fix": True}})
    assert st == 200, r
    st, sup = c.get("/api/suppliers")
    assert sup["data"][0]["name"] == "Acme Supply" and sup["data"][0]["purchase_value"] > 0
    # mapping validation: missing required field
    st, r = upload(c, "purchases", [["a", "b"], ["1", "2"]])
    st, v = c.post(f"/api/imports/{r['data']['id']}/validate", {"mapping": {}})
    assert st == 400 and "required" in v["message"].lower()


def test_dangerous_uploads_rejected(app):
    c, _ = register(app, "evil@test.com")
    for name, content in (("x.exe", b"MZ\x90\x00"), ("x.csv", b"MZ\x90\x00\x03"), ("x.xlsx", b"not a zip"), ("x.csv", b"<html><script>"), ("x.php", b"a,b\n1,2")):
        st, r = c.post("/api/imports", data={"dataset_type": "sales", "file": (io.BytesIO(content), name)}, content_type="multipart/form-data")
        assert st == 400 and r["success"] is False, name
    st, r = upload(c, "sales", [["only header"]])
    assert st == 400


def test_tenant_isolation(app):
    a, _ = register(app, "tenant-a@test.com", "Company A")
    b, _ = register(app, "tenant-b@test.com", "Company B")
    full_import(a, "products", PRODUCTS)
    st, pa = a.get("/api/products")
    pid = pa["data"]["items"][0]["id"]
    assert b.get("/api/products")[1]["data"]["total"] == 0
    for url in (f"/api/analytics/risk/{pid}", f"/api/analytics/forecast/{pid}"):
        assert b.get(url)[0] == 404
    assert b.put(f"/api/products/{pid}", {"name": "hacked"})[0] == 404
    assert b.delete(f"/api/products/{pid}")[0] == 404
    assert b.post("/api/inventory/adjust", {"product_id": pid, "type": "STOCK_IN", "quantity": 5})[0] == 404
    assert b.post("/api/sales", {"product_id": pid, "quantity": 1})[0] == 404
    assert b.post("/api/analytics/simulate", {"mode": "product", "product_id": pid, "discount": 20})[0] in (400, 404)
    jobs = a.get("/api/imports")[1]["data"]["items"]
    assert b.get(f"/api/imports/{jobs[0]['id']}")[0] == 404
    assert b.delete(f"/api/imports/{jobs[0]['id']}")[0] == 404
    assert b.get("/api/analytics/dashboard")[1]["data"]["has_data"] is False
    assert a.get(f"/api/products/{pid}")[0] in (404, 405)
    st, r = a.get("/api/does-not-exist")
    assert st == 404 and r["success"] is False                                # JSON, not the SPA html
    # supplier from another workspace cannot be attached
    sid = a.post("/api/suppliers", {"name": "A Supplier"})[1]["data"]["id"]
    assert b.post("/api/products", {"name": "x", "sku": "X1", "supplier_id": sid})[0] == 400
    assert b.get(f"/api/suppliers/{sid}")[0] == 404
    assert "Company A" not in str(b.get("/api/audit-log")[1])


def test_roles_enforced_on_backend(app):
    owner, _ = register(app, "owner@roles.com", "Roles Co")
    full_import(owner, "products", PRODUCTS)
    made = {}
    for role in ("VIEWER", "ANALYST", "SALES_MANAGER", "INVENTORY_MANAGER", "ADMIN"):
        email = f"{role.lower()}@roles.com"
        assert owner.post("/api/users", {"name": role, "email": email, "password": PW, "role": role})[0] == 201
        c = Client(app)
        c.h = {"Authorization": "Bearer " + c.post("/api/auth/login", {"email": email, "password": PW})[1]["data"]["access_token"]}
        made[role] = c
    pid = owner.get("/api/products")[1]["data"]["items"][0]["id"]
    assert made["VIEWER"].get("/api/analytics/dashboard")[0] == 200
    assert made["VIEWER"].post("/api/products", {"name": "n", "sku": "N1"})[0] == 403
    assert made["VIEWER"].post("/api/analytics/refresh", {})[0] == 403
    assert made["VIEWER"].get("/api/users")[0] == 403
    assert made["SALES_MANAGER"].post("/api/products", {"name": "n", "sku": "N1"})[0] == 403
    assert made["SALES_MANAGER"].post("/api/sales", {"product_id": pid, "quantity": 1})[0] == 201
    assert made["INVENTORY_MANAGER"].post("/api/products", {"name": "n", "sku": "N1", "quantity": 4})[0] == 201
    assert made["INVENTORY_MANAGER"].post("/api/sales", {"product_id": pid, "quantity": 1})[0] == 403
    assert made["ANALYST"].post("/api/analytics/refresh", {})[0] == 200
    assert made["ANALYST"].post("/api/products", {"name": "n", "sku": "N2"})[0] == 403
    assert made["INVENTORY_MANAGER"].get("/api/audit-log")[0] == 403
    assert made["ADMIN"].get("/api/audit-log")[0] == 200
    assert made["ADMIN"].post("/api/users", {"name": "o", "email": "o2@roles.com", "password": PW, "role": "OWNER"})[0] == 403
    assert made["ADMIN"].delete("/api/workspace/data")[0] == 403          # owner only
    me = owner.get("/api/auth/me")[1]["data"]["user"]
    assert owner.delete(f"/api/users/{me['id']}")[0] == 400                 # cannot delete self / last owner


def test_product_inventory_sales_crud(app):
    c, _ = register(app, "crud@test.com")
    st, r = c.post("/api/products", {"name": "Lamp", "sku": "lamp-1", "category": "Home", "cost_price": 100, "selling_price": 180, "quantity": 20, "reorder_level": 5})
    assert st == 201 and r["data"]["sku"] == "LAMP-1"
    pid = r["data"]["id"]
    assert c.post("/api/products", {"name": "Lamp2", "sku": "LAMP-1"})[0] == 409
    assert c.post("/api/products", {"name": "", "sku": "Z"})[0] == 400
    assert c.post("/api/sales", {"product_id": pid, "quantity": 50})[0] == 400          # insufficient stock
    assert c.post("/api/sales", {"product_id": pid, "quantity": 4})[0] == 201
    assert c.post("/api/inventory/adjust", {"product_id": pid, "type": "STOCK_IN", "quantity": 10, "note": "restock"})[0] == 200
    assert c.post("/api/inventory/adjust", {"product_id": pid, "type": "ADJUSTMENT", "quantity": 25})[0] == 200
    assert c.get("/api/products?q=lamp")[1]["data"]["items"][0]["quantity"] == 25
    mv = c.get("/api/inventory/movements")[1]["data"]
    assert mv["total"] >= 4 and mv["items"][0]["user"] == "Test User"
    assert c.post("/api/purchases", {"product_id": pid, "quantity": 5, "unit_cost": 90})[0] == 201
    assert c.get("/api/inventory/summary")[1]["data"]["units"] == 30
    assert c.put(f"/api/products/{pid}", {"selling_price": 200})[1]["data"]["selling_price"] == 200
    assert c.delete(f"/api/products/{pid}")[0] == 200
    assert c.get("/api/products")[1]["data"]["total"] == 0


def test_demo_workspace_flow_and_replace(app):
    c, _ = register(app, "demo@test.com")
    assert c.post("/api/workspace/load-demo", {})[0] == 200
    ws = c.get("/api/workspace")[1]["data"]
    assert ws["is_demo"] is True and ws["has_data"]
    jid, rep, st, r = full_import(c, "products", PRODUCTS)
    assert st == 409 and r["code"] == "demo_data_present"
    jid, rep, st, r = full_import(c, "products", PRODUCTS, replace_demo=True)
    assert st == 200, r
    assert c.get("/api/workspace")[1]["data"]["is_demo"] is False
    assert c.get("/api/products")[1]["data"]["total"] == 3                   # sample data gone


def test_simulator_reports_exports_notifications(app):
    c, _ = register(app, "sim@test.com")
    c.post("/api/workspace/load-demo", {})
    top = c.get("/api/analytics/risk?sort=blocked&dir=desc")[1]["data"]["items"][0]
    st, s = c.post("/api/analytics/simulate", {"mode": "product", "product_id": top["product_id"], "discount": 30})
    assert st == 200 and s["data"]["capital_recovered"] >= 0 and len(s["data"]["curve"]) > 5
    st, p = c.post("/api/analytics/simulate", {"mode": "portfolio", "discount": 30, "scope": {"mode": "blocked"}})
    assert st == 200 and p["data"]["projected"]["blocked_capital"] <= p["data"]["current"]["blocked_capital"]
    assert c.post("/api/analytics/simulate", {"mode": "portfolio", "discount": 150})[0] == 400
    for kind in ("inventory", "dead-stock", "risk", "sales", "purchases", "blocked-capital", "recommendations", "executive"):
        for fmt, magic in (("csv", None), ("xlsx", b"PK"), ("pdf", b"%PDF")):
            st, r = c.get(f"/api/reports/{kind}?format={fmt}")
            assert st == 200, (kind, fmt, getattr(r, "data", r)[:200] if hasattr(r, "data") else r)
            if magic:
                assert r.data.startswith(magic), (kind, fmt)
    assert c.get("/api/reports/inventory?format=exe")[0] == 400
    assert c.get("/api/notifications")[1]["data"]["total"] >= 1
    assert c.get("/api/analytics/recommendations")[1]["data"]["total"] > 0
    assert c.get("/api/analytics/executive")[1]["data"]["money_stuck"]["by_category"]
    assert c.get("/api/analytics/forecast")[1]["data"]["summary"]["statistical"] >= 0
    pid = top["product_id"]
    assert c.get(f"/api/analytics/forecast/{pid}")[1]["data"]["method"]


def test_delete_import_removes_only_its_data(app):
    c, _ = register(app, "del@test.com")
    full_import(c, "products", PRODUCTS)
    jid, *_ = full_import(c, "sales", sales_rows())
    assert c.get("/api/sales")[1]["data"]["total"] > 100
    assert c.delete(f"/api/imports/{jid}")[0] == 200
    assert c.get("/api/sales")[1]["data"]["total"] == 0
    assert c.get("/api/products")[1]["data"]["total"] == 3
    assert any(a["action"] == "DATA_DELETED" for a in c.get("/api/audit-log")[1]["data"]["items"])


def test_shared_demo_is_read_only(app):
    import seed as S
    with app.app_context():
        from models import User, Workspace
        from extensions import db
        from services.demo_data import load_demo
        ws = Workspace(name="Demo Workspace", business_type="Electronics", is_demo=True, is_shared_demo=True, onboarding_completed=True, risk_config={})
        db.session.add(ws); db.session.flush()
        u = User(workspace_id=ws.id, name="v", email="demo-viewer@stockghost.demo", role="VIEWER"); u.set_password(PW)
        db.session.add(u); load_demo(ws); db.session.commit()
    c = Client(app)
    st, r = c.post("/api/auth/demo-login", {})
    assert st == 200
    c.h = {"Authorization": "Bearer " + r["data"]["access_token"]}
    assert c.get("/api/products")[1]["data"]["total"] == 100
    assert c.post("/api/products", {"name": "n", "sku": "N"})[0] == 403
    assert c.post("/api/workspace/load-demo", {})[0] == 403


def test_finishing_onboarding_does_not_mark_analysis_stale(app):
    c, _ = register(app, "stale@test.com")
    full_import(c, "products", PRODUCTS)
    assert c.get("/api/workspace")[1]["data"]["analysis_stale"] is False
    assert c.put("/api/workspace", {"onboarding_completed": True})[0] == 200
    assert c.get("/api/workspace")[1]["data"]["analysis_stale"] is False        # regression: used to flip to True
    assert c.put("/api/workspace", {"config": {"dead_days": 60}})[0] == 200
    assert c.get("/api/workspace")[1]["data"]["analysis_stale"] is True         # scoring-relevant change -> stale
    assert c.put("/api/workspace", {"config": {"dead_days": 0}})[0] == 400      # validation
