"""Report builders (data) and exporters (CSV / Excel / PDF). Everything is scoped to one workspace."""
import csv
import datetime as dt
import io

from sqlalchemy import func
from sqlalchemy.orm import joinedload

from models import Product, Purchase, RiskAnalysis, Sale, Supplier
from services.analytics import dashboard
from services.dsrs import LEVEL_LABELS, get_config
from services.import_fields import DATASETS  # noqa: F401  (kept for symmetry)
from services.importer import safe_cell
from services.recommendations import all_recommendations
from utils import inr

TITLES = {"inventory": "Inventory Report", "dead-stock": "Dead Stock Report", "risk": "Risk Report", "sales": "Sales Report",
          "purchases": "Purchase Report", "blocked-capital": "Blocked Capital Report", "recommendations": "Recommendation Report",
          "executive": "Executive Summary"}
MAX_ROWS = {"json": 500, "csv": 100000, "xlsx": 50000, "pdf": 40}


def _risk_rows(ws_id):
    return (RiskAnalysis.query.filter(RiskAnalysis.workspace_id == ws_id).join(Product, Product.id == RiskAnalysis.product_id)
            .options(joinedload(RiskAnalysis.product).joinedload(Product.supplier)))


def build(kind, ws, start=None, end=None, limit=100000):
    wid, today = ws.id, dt.date.today()
    meta = {"title": TITLES[kind], "company": ws.name, "demo": ws.is_demo,
            "period": f"{start} to {end}" if start and end else f"As of {today.strftime('%d %b %Y')}",
            "generated_at": dt.datetime.now().strftime("%d %b %Y, %H:%M")}
    chart = None
    if kind == "risk":
        rs = _risk_rows(wid).order_by(RiskAnalysis.dsrs_score.desc(), RiskAnalysis.blocked_capital.desc()).limit(limit).all()
        cols = ["SKU", "Product", "Category", "Quantity", "Stock age (days)", "Days since sale", "Inventory value", "DSRS", "Risk", "Status", "Blocked capital", "Why"]
        rows = [[r.product.sku, r.product.name, r.product.category, r.product.quantity, None if r.stock_age is None else round(r.stock_age),
                 r.last_sale_days, r.inventory_value, r.dsrs_score, LEVEL_LABELS[r.risk_level], r.stock_status, r.blocked_capital,
                 "; ".join((r.details or {}).get("reasons", [])[:3])] for r in rs]
        totals = {"Products": len(rows), "Inventory value": round(sum(r[6] for r in rows), 2), "Blocked capital": round(sum(r[10] for r in rows), 2)}
        counts = {}
        for r in rs:
            counts[LEVEL_LABELS[r.risk_level]] = counts.get(LEVEL_LABELS[r.risk_level], 0) + 1
        chart = {"title": "Products by risk level", "labels": list(counts), "values": list(counts.values()), "money": False}
    elif kind == "dead-stock":
        rs = (_risk_rows(wid).filter(RiskAnalysis.stock_status.in_(("Dead Stock", "Slow Moving", "Overstock")))
              .order_by(RiskAnalysis.inventory_value.desc()).limit(limit).all())
        cols = ["SKU", "Product", "Category", "Status", "Quantity", "Inventory value", "Days since sale", "Sales velocity /day", "DSRS", "Blocked capital"]
        rows = [[r.product.sku, r.product.name, r.product.category, r.stock_status, r.product.quantity, r.inventory_value, r.last_sale_days,
                 r.avg_daily_demand, r.dsrs_score, r.blocked_capital] for r in rs]
        totals = {"Products": len(rows), "Inventory value": round(sum(r[5] for r in rows), 2), "Blocked capital": round(sum(r[9] for r in rows), 2)}
        agg = {}
        for r in rs:
            agg[r.stock_status] = agg.get(r.stock_status, 0) + r.inventory_value
        chart = {"title": "Inventory value by status", "labels": list(agg), "values": [round(v) for v in agg.values()], "money": True}
    elif kind == "blocked-capital":
        rs = _risk_rows(wid).filter(RiskAnalysis.blocked_capital > 0).order_by(RiskAnalysis.blocked_capital.desc()).limit(limit).all()
        cols = ["SKU", "Product", "Category", "Supplier", "Warehouse", "Quantity", "Cost price", "Blocked capital", "DSRS", "Risk"]
        rows = [[r.product.sku, r.product.name, r.product.category, r.product.supplier.name if r.product.supplier else "", r.product.warehouse or "",
                 r.product.quantity, r.product.cost_price, r.blocked_capital, r.dsrs_score, LEVEL_LABELS[r.risk_level]] for r in rs]
        totals = {"Blocked products": len(rows), "Blocked capital": round(sum(r[7] for r in rows), 2)}
        agg = {}
        for r in rs:
            agg[r.product.category] = agg.get(r.product.category, 0) + r.blocked_capital
        top = sorted(agg.items(), key=lambda kv: -kv[1])[:8]
        chart = {"title": "Blocked capital by category", "labels": [k for k, _ in top], "values": [round(v) for _, v in top], "money": True}
    elif kind == "inventory":
        rs = _risk_rows(wid).order_by(Product.category, Product.name).limit(limit).all()
        cols = ["SKU", "Product", "Category", "Supplier", "Warehouse", "Cost price", "Selling price", "Quantity", "Reorder level", "Inventory value", "Status"]
        rows = [[r.product.sku, r.product.name, r.product.category, r.product.supplier.name if r.product.supplier else "", r.product.warehouse or "",
                 r.product.cost_price, r.product.selling_price, r.product.quantity, r.product.reorder_level, r.inventory_value, r.stock_status] for r in rs]
        totals = {"Products": len(rows), "Units": sum(r[7] for r in rows), "Inventory value": round(sum(r[9] for r in rows), 2)}
        agg = {}
        for r in rs:
            agg[r.product.category] = agg.get(r.product.category, 0) + r.inventory_value
        top = sorted(agg.items(), key=lambda kv: -kv[1])[:8]
        chart = {"title": "Inventory value by category", "labels": [k for k, _ in top], "values": [round(v) for _, v in top], "money": True}
    elif kind == "sales":
        q = (Sale.query.filter(Sale.workspace_id == wid).options(joinedload(Sale.product)))
        if start:
            q = q.filter(Sale.sale_date >= start)
        if end:
            q = q.filter(Sale.sale_date <= end)
        ss = q.order_by(Sale.sale_date.desc(), Sale.id.desc()).limit(limit).all()
        cols = ["Date", "SKU", "Product", "Category", "Quantity", "Unit price", "Revenue", "Customer", "Region"]
        rows = [[s.sale_date.isoformat(), s.product.sku, s.product.name, s.product.category, s.quantity, s.unit_price or 0,
                 round(s.quantity * (s.unit_price or 0), 2), s.customer or "", s.region or ""] for s in ss]
        totals = {"Sales shown": len(rows), "Units": sum(r[4] for r in rows), "Revenue": round(sum(r[6] for r in rows), 2)}
        agg = {}
        for r in rows:
            agg[r[0][:7]] = agg.get(r[0][:7], 0) + r[6]
        keys = sorted(agg)[-12:]
        chart = {"title": "Revenue by month", "labels": keys, "values": [round(agg[k]) for k in keys], "money": True}
    elif kind == "purchases":
        q = Purchase.query.filter(Purchase.workspace_id == wid).options(joinedload(Purchase.product), joinedload(Purchase.supplier))
        if start:
            q = q.filter(Purchase.purchase_date >= start)
        if end:
            q = q.filter(Purchase.purchase_date <= end)
        ps = q.order_by(Purchase.purchase_date.desc(), Purchase.id.desc()).limit(limit).all()
        cols = ["Date", "SKU", "Product", "Supplier", "Quantity", "Unit cost", "Total cost", "Reference"]
        rows = [[p.purchase_date.isoformat(), p.product.sku, p.product.name, p.supplier.name if p.supplier else "", p.quantity,
                 p.unit_cost if p.unit_cost is not None else p.product.cost_price, p.to_dict()["total_cost"], p.reference or ""] for p in ps]
        totals = {"Purchases shown": len(rows), "Units": sum(r[4] for r in rows), "Total cost": round(sum(r[6] for r in rows), 2)}
        agg = {}
        for r in rows:
            agg[r[0][:7]] = agg.get(r[0][:7], 0) + r[6]
        keys = sorted(agg)[-12:]
        chart = {"title": "Purchase cost by month", "labels": keys, "values": [round(agg[k]) for k in keys], "money": True}
    elif kind == "recommendations":
        recs = all_recommendations(ws)[:limit]
        cols = ["SKU", "Product", "Priority", "DSRS", "Recommended action", "Suggested discount %", "Blocked capital", "Estimated capital recovery (estimate)",
                "Estimated units reduced (estimate)", "Why"]
        rows = [[r["sku"] or "", r["product_name"], r["priority"], r["dsrs_score"], r["action"], r["suggested_discount"], r["blocked_capital"],
                 (r["estimate"] or {}).get("capital_recovery", ""), (r["estimate"] or {}).get("units_reduction", ""), " ".join(r["reasons"][:3])] for r in recs]
        totals = {"Actions": len(rows), "Estimated capital recovery": round(sum(r[7] for r in rows if isinstance(r[7], (int, float))), 2)}
        agg = {}
        for r in rows:
            agg[r[2]] = agg.get(r[2], 0) + 1
        chart = {"title": "Actions by priority", "labels": list(agg), "values": list(agg.values()), "money": False}
    else:                                                          # executive
        d = dashboard_for_report(ws)
        k = d.get("kpis", {})
        cols = ["Metric", "Value"]
        rows = [["Total inventory value", k.get("inventory_value")], ["Blocked capital", k.get("blocked_capital")], ["Blocked % of inventory", k.get("blocked_pct")],
                ["Dead stock value", k.get("dead_stock_value")], ["Dead stock products", k.get("dead_stock_count")], ["High / critical risk products", k.get("high_risk")],
                ["Slow-moving products", k.get("slow_moving")], ["Stockout-risk products", k.get("stockout_risk")], ["Inventory turnover (annualised)", k.get("inventory_turnover")],
                ["Sales revenue (last 30 days)", k.get("sales_revenue")], ["Potential recovery (estimate)", k.get("potential_recovery")], ["Average DSRS", k.get("average_dsrs")]]
        totals = {}
        by = d.get("blocked_by_category", [])[:8]
        chart = {"title": "Blocked capital by category", "labels": [b["category"] for b in by], "values": [round(b["blocked_capital"]) for b in by], "money": True}
        meta["kpis"] = k
    return {"meta": meta, "columns": cols, "rows": rows, "totals": totals, "chart": chart}


def dashboard_for_report(ws):
    from flask import current_app
    with current_app.test_request_context("/?range=30d"):
        return dashboard(ws)


# ---- exporters -------------------------------------------------------------------------------------------------------
def to_csv(rep):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(rep["columns"])
    for r in rep["rows"]:
        w.writerow([safe_cell(c) if isinstance(c, str) else ("" if c is None else c) for c in r])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(rep):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"[:31]
    head = PatternFill("solid", fgColor="0F1626")
    ws.append([rep["meta"]["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{rep['meta']['company']}{' (DEMO DATA)' if rep['meta']['demo'] else ''} | {rep['meta']['period']} | Generated {rep['meta']['generated_at']}"])
    ws.append([])
    ws.append(rep["columns"])
    hr = ws.max_row
    for c in ws[hr]:
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), head, Alignment(horizontal="center", wrap_text=True)
    for r in rep["rows"]:
        ws.append([safe_cell(c) if isinstance(c, str) else c for c in r])
    for i, col in enumerate(rep["columns"], start=1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in rep["rows"][:300] if i - 1 < len(r)]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(48, width)
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    if rep["totals"]:
        t = wb.create_sheet("Summary")
        t.append(["Metric", "Value"])
        for k, v in rep["totals"].items():
            t.append([k, v])
        t.column_dimensions["A"].width = 32
        t.column_dimensions["B"].width = 20
        for c in t[1]:
            c.font = Font(bold=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _rs(v):
    """PDF-safe money (standard PDF fonts have no rupee glyph)."""
    return inr(v).replace("₹", "Rs ")


def to_pdf(rep, ws):
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    buf = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=rep["meta"]["title"], author="StockGhost")
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontSize=20, alignment=0, textColor=colors.HexColor("#0b1120"), spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12, textColor=colors.HexColor("#0f766e"), spaceBefore=10, spaceAfter=4)
    sm = ParagraphStyle("sm", parent=ss["BodyText"], fontSize=8, leading=10, textColor=colors.HexColor("#475569"))
    cell = ParagraphStyle("cell", parent=ss["BodyText"], fontSize=7, leading=8.5)
    meta = rep["meta"]
    el = [Paragraph(meta["title"], h1),
          Paragraph(f"<b>{_esc(meta['company'])}</b>{' - DEMO DATA' if meta['demo'] else ''} &nbsp;|&nbsp; Period: {meta['period']} &nbsp;|&nbsp; Generated: {meta['generated_at']}", sm),
          Spacer(1, 6)]

    def money_cell(c, col):
        if isinstance(c, float) and any(w in col.lower() for w in ("value", "capital", "price", "cost", "revenue", "recovery", "sales")):
            return _rs(c)
        return "" if c is None else str(c)

    # KPIs
    kp = []
    for k, v in (rep["totals"] or {}).items():
        kp.append([k, _rs(v) if isinstance(v, float) and any(w in k.lower() for w in ("value", "capital", "cost", "revenue", "recovery")) else f"{v:,}" if isinstance(v, int) else str(v)])
    if rep["meta"].get("kpis"):
        kk = rep["meta"]["kpis"]
        kp = [["Inventory value", _rs(kk["inventory_value"])], ["Blocked capital", f"{_rs(kk['blocked_capital'])} ({kk['blocked_pct']}%)"],
              ["Dead stock value", _rs(kk["dead_stock_value"])], ["High / critical risk products", str(kk["high_risk"])],
              ["Stockout-risk products", str(kk["stockout_risk"])], ["Revenue (30 days)", _rs(kk["sales_revenue"])],
              ["Potential recovery (estimate)", _rs(kk["potential_recovery"])]]
    if kp:
        el.append(Paragraph("Key figures", h2))
        t = Table([kp[i:i + 4] and sum(([a, b] for a, b in kp[i:i + 4]), []) for i in range(0, len(kp), 4)], colWidths=[48 * mm, 36 * mm] * 4)
        t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 8), ("FONTNAME", (0, 0), (-1, -1), "Helvetica"), ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#334155")),
                               ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f1f5f9")), ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                               ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0"))] +
                              [("FONTNAME", (c, 0), (c, -1), "Helvetica-Bold") for c in (1, 3, 5, 7)]))
        el.append(t)
    ch = rep.get("chart")
    if ch and ch["labels"]:
        el.append(Paragraph(ch["title"], h2))
        d = Drawing(240 * mm, 62 * mm)
        bc = VerticalBarChart()
        bc.x, bc.y, bc.width, bc.height = 14 * mm, 14 * mm, 215 * mm, 42 * mm
        bc.data = [ch["values"]]
        bc.categoryAxis.categoryNames = [str(x)[:16] for x in ch["labels"]]
        bc.categoryAxis.labels.fontSize = 7
        bc.valueAxis.labels.fontSize = 7
        bc.valueAxis.valueMin = 0
        bc.bars[0].fillColor = colors.HexColor("#14b8a6")
        bc.valueAxis.labelTextFormat = (lambda v: f"{v / 1e5:.1f}L" if v >= 1e5 else f"{v:,.0f}") if ch["money"] else "%d"
        d.add(bc)
        if ch["money"]:
            d.add(String(14 * mm, 2 * mm, "Values in Rs (L = lakh)", fontSize=7, fillColor=colors.HexColor("#64748b")))
        el.append(d)

    def table(cols, rows, width_hint=None):
        data = [[Paragraph(f"<b>{_esc(c)}</b>", cell) for c in cols]] + [[Paragraph(_esc(money_cell(c, cols[i])), cell) for i, c in enumerate(r)] for r in rows]
        tb = Table(data, repeatRows=1, colWidths=[(page[0] - 28 * mm) / len(cols)] * len(cols))
        tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (1, 0), (-1, -1), [colors.white, colors.HexColor("#f8fafc")])]))
        return tb

    if rep["meta"]["title"] != TITLES["executive"]:
        shown = rep["rows"][:MAX_ROWS["pdf"]]
        el.append(Paragraph(f"{meta['title']} (showing {len(shown)} of {len(rep['rows'])} rows - use CSV/Excel for the full list)", h2))
        el.append(table(rep["columns"], shown))
    else:
        el.append(Paragraph("Key metrics", h2))
        el.append(table(rep["columns"], rep["rows"]))
    top = _risk_rows(ws.id).order_by(RiskAnalysis.dsrs_score.desc(), RiskAnalysis.blocked_capital.desc()).limit(10).all()
    if top:
        el.append(Paragraph("Top risk products", h2))
        el.append(table(["Product", "SKU", "DSRS", "Risk", "Status", "Blocked capital", "Main reason"],
                        [[r.product.name, r.product.sku, r.dsrs_score, LEVEL_LABELS[r.risk_level], r.stock_status, r.blocked_capital,
                          (r.details or {}).get("reasons", [""])[0] if (r.details or {}).get("reasons") else ""] for r in top]))
    recs = all_recommendations(ws)[:8]
    if recs:
        el.append(Paragraph("Recommended actions (estimates)", h2))
        el.append(table(["Product", "Priority", "Action", "Why", "Est. capital recovery"],
                        [[r["product_name"], r["priority"], r["action"], " ".join(r["reasons"][:2]), (r["estimate"] or {}).get("capital_recovery", "")] for r in recs]))
    el.append(Spacer(1, 8))
    el.append(Paragraph("DSRS is a rule-based risk score; forecasts are statistical (not machine learning); recovery figures are projections based on stated assumptions, not guarantees.", sm))

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(14 * mm, 8 * mm, f"StockGhost - {meta['company']}")
        canvas.drawRightString(page[0] - 14 * mm, 8 * mm, f"Page {d.page}")
        canvas.restoreState()
    doc.build(el, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
